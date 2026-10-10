"""Non-destructive editorial revision and editable Word export from sealed TeX.

No simulation, bootstrap, or raw experiment inputs. Existing vector figures
are byte-reused, with 300 dpi previews embedded in Word (not replacement data).
"""
from pathlib import Path
import hashlib
import json
import re
import shutil
import subprocess
import zipfile
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'manuscript/joconline-integrated-v1'
OUT = ROOT / 'manuscript/joconline-submission-v1'
QA = ROOT / 'results/diagnostics/joconline-submission-20261010'
DEPS = Path('C:/Users/jiate/.cache/codex-runtimes/codex-primary-runtime/dependencies')
PANDOC = ROOT / 'outputs/merged_manuscript_latex_20260918/_tools/pandoc/pandoc-3.11/pandoc.exe'
POPPLER = DEPS / 'native/poppler/Library/bin/pdftoppm.exe'

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def write_json(p, obj):
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2)+'\n', encoding='utf-8', newline='\n')

def replace_once(s, old, new):
    assert s.count(old) == 1, old[:70]
    return s.replace(old, new, 1)

def revise():
    assert not OUT.exists(), 'Preserve an existing delivery; choose a new version.'
    manifest = json.loads((BASE/'manifest.json').read_text('utf-8'))
    for item in manifest['files']:
        p = ROOT/item['path']
        assert p.stat().st_size == item['bytes'] and sha(p) == item['sha256'], p
    OUT.mkdir(); QA.mkdir(parents=True, exist_ok=True)
    for directory in ('figures', 'source-data'):
        shutil.copytree(BASE/directory, OUT/directory)
    for name in ('.gitattributes', '.gitignore', 'supplement.tex', 'source-validation.json', 'claim-evidence-map.md'):
        shutil.copyfile(BASE/name, OUT/name)
    main = (BASE/'main.tex').read_text('utf-8')
    original = main
    start = main.index('摘要：')+3; end = main.index('\\par}', start)
    abstract = ('针对超图支付网络持续服务评价中的资源口径差异，建立资金与参与关系匹配的评估方法。'
                '以首次无可用路径为服务事件，结合原子支付、余额感知路由和父图分层重抽样，比较4类超图构造与二元参照。'
                '两独立相位各含240个父图区块，40项比较均满足复现准则；确认相位总体失败风险差为$-$0.449～$-$0.787。'
                '所设合成场景中的服务收益伴随更高协调暴露量；事后补充未支持搜索的校正家族增益，优化初态亦非普遍有利。')
    assert len(abstract.replace('$-$','-')) <= 200
    main = main[:start]+abstract+main[end:]
    start = main.index('Abstract: ')+10; end = main.index('\\par}', start)
    main = main[:start]+('A resource-matched evaluation method was developed to assess sustained service in hypergraph payment networks. '
        'Funding and incidence budgets were controlled. First path unavailability defined the service event. '
        'Atomic payments, balance-aware routing and parent-stratified bootstrap resampling were used to compare four hypergraph construction families with binary references. '
        'Two independent phases each comprised 240 parent-graph blocks; all 40 contrasts met the registered replication criterion. '
        'Confirmation-phase global failure-risk differences ranged from $-$0.449 to $-$0.787. '
        'Under the specified synthetic conditions, service benefits co-occur with higher coordination exposure. '
        'Post-hoc analyses do not support family-adjusted search gains, and optimized initialization is not uniformly beneficial.')+main[end:]
    main = replace_once(main,
        '首先，将单超边内的守恒状态推广为多超边关联坐标，明确局部余额耗尽、全局无路径和协议拒绝的区别，并说明路由选择如何引入跨超边依赖。',
        '首先，在已有多方信道余额模型基础上，统一局部余额耗尽、全局无路径和协议拒绝的记录口径，明确单超边边界事件与多超边请求服务事件的区别；不将关联坐标表示或守恒更新本身作为新模型。')
    main = replace_once(main,
        '文献\\upcite{r10}和文献\\upcite{r11}分别提供了超图构造与多方路径规划的直接基础。本研究沿用相关构造类别，明确闭邻域节点覆盖规则，并增加需求感知局部搜索、二元参与关系匹配和跨相位确认。已有数学模型讨论了多方信道的可行状态与流动性几何结构\\upcite{r14}；Starfish探索基于局部星形结构的多方重平衡，其修订预印本已注明被High-Confidence Computing录用\\upcite{r15}。这些研究与本工作互补。本文的贡献定位于受控比较方法及其复现证据。',
        '文献\\upcite{r10}和文献\\upcite{r11}分别提供超图构造与多方路径规划的直接基础，已有研究已采用成员余额表示及跨信道局部转移，并评价支付成功率和成本。本文沿用相关构造类别与余额模型，重点检验受控资源下的序列服务差异。已有数学模型讨论多方信道的可行状态与流动性几何结构\\upcite{r14}；Starfish探索基于局部星形结构的多方重平衡，其修订预印本注明被High-Confidence Computing录用，并已登记期刊DOI\\upcite{r15}。这些工作与本文的比较评估互补。\n\n'
        '近期预印本COALESCE以DAG状态承诺和转移证明组织跨超边支付，关注多方协议执行与协调，并报告成功率、耗尽及路径长度\\upcite{r16}。MaxPTE通过支付拓扑熵指导信道合并，在保持资金总量的同时重配置结构；其比较控制总信道数，并以当前余额下的最大流不足定义失败比例\\upcite{r17}。后者不同于本文原子单路径请求序列中的首次无路径事件。因此，本文不声称首次研究多超边支付或首次采用资源控制，也不将不同协议与终点条件下的文献数值用于性能排名；贡献限定于节点资金和参与关系匹配、父图配对分层推断及双相位复现的组合评估。')
    main = replace_once(main, '设支付超图为$\\mathcal H=(V,\\mathcal E)$，',
        '沿用超图支付与多方路径规划的成员余额表示\\upcite{r10,r11}，设支付超图为$\\mathcal H=(V,\\mathcal E)$，')
    main = replace_once(main, '2025, 4(4): 1-14. DOI: 10.1145/3702248.',
        '2025, 4(4): 30:1-30:14. DOI: 10.1145/3702248.')
    main = replace_once(main, '\\url{https://arxiv.org/abs/2504.20536v2}. 修订预印本.',
        '\\url{https://arxiv.org/abs/2504.20536v2}. 修订预印本；已登记期刊DOI: 10.1016/j.hcc.2026.100443.')
    extra = ('\\bibitem{r16} NAINWAL A, KAMBLE A, AWATHARE N. Hypergraph based multi-party payment channel[EB/OL]. (2026-06-02)[2026-10-10]. \\url{https://arxiv.org/abs/2512.11775v2}. 修订预印本.\n\n'
             '\\bibitem{r17} XIAO S, WANG S, SHI H, et al. Employing the structural power to achieve supply-demand balanced payment channel networks[EB/OL]. (2026-09-03)[2026-10-10]. \\url{https://arxiv.org/abs/2609.03600v1}. 预印本.\n\n')
    main = main.replace('\\end{thebibliography}', extra+'\\end{thebibliography}', 1)
    main = main.replace(r'\begin{thebibliography}{99}\small\raggedright\setlength{\itemsep}{2pt}',
        r'\begin{thebibliography}{99}\small\linespread{0.96}\selectfont\raggedright\setlength{\itemsep}{0pt}')
    (OUT/'main.tex').write_text(main, encoding='utf-8', newline='\n')
    equations = lambda t: re.findall(r'\\begin\{equation\}(.*?)\\end\{equation\}', t, re.S)
    assert equations(main) == equations(original) and len(equations(main)) == 18
    first = []
    for group in re.findall(r'\\upcite\{([^}]+)\}', main):
        for key in group.split(','):
            if key not in first: first.append(key)
    assert first == [f'r{i}' for i in range(1,18)]
    write_json(OUT/'revision-checks.json', dict(sealed_manifest_sha256=sha(BASE/'manifest.json'),
        sealed_files_verified=len(manifest['files']), main_source_sha256=sha(BASE/'main.tex'),
        displayed_equations_unchanged=18, references=17, first_citation_order=first,
        chinese_abstract_characters=len(abstract.replace('$-$','-')), new_simulations=0,
        new_bootstrap_replicates=0, scientific_figures_byte_reused=11))

def prepare_tex(name):
    s = (OUT/(name+'.tex')).read_text('utf-8').split('\\begin{document}',1)[1].split('\\end{document}',1)[0]
    s = s.replace('\\twocolumn[{\\begin{minipage}{\\textwidth}', '').replace('\\vspace{7pt}\\end{minipage}}]', '')
    for kind in ('figure','table'):
        s = s.replace(r'\renewcommand{\the'+kind+r'}{S\arabic{'+kind+'}}','')
    s = s.replace('\\onecolumn','').replace('\\balance','')
    # Word math converter requires modern spelling of the same upright labels.
    s = re.sub(r'\\rm\s+(same|shift)',lambda m:r'\mathrm{'+m[1]+'}',s)
    # Pandoc does not implement keepaspectratio when both dimensions are given.
    s = re.sub(r',height=\d+mm,keepaspectratio', '', s)
    fig_no=[0]; tab_no=[0]; kind=['table']; prefix='S' if name=='supplement' else ''
    def caption(m):
        if m[1]:kind[0]=m[1].rstrip('*');return m[0]
        counter=fig_no if kind[0]=='figure' else tab_no
        counter[0]+=1
        return r'\caption{'+('图' if kind[0]=='figure' else '表')+prefix+str(counter[0])+' '+m[2]+'}'
    s=re.sub(r'\\begin\{(figure\*?|table\*?)\}|\\caption\{([^}]*)\}',caption,s)
    def outside_caption(m):
        fig_no[0]+=1
        return r'\par\noindent 图S'+str(fig_no[0])+' '+m[1]+r'\par'
    s=re.sub(r'\\captionof\{figure\}\{([^}]*)\}',outside_caption,s)
    labels = {}
    for kind in ('eq','fig','tab'):
        for number,label in enumerate(re.findall(r'\\label\{('+kind+r':[^}]+)\}',s),1):
            labels[label] = str(number)
    s = re.sub(r'\\ref\{([^}]+)\}',lambda m:labels[m[1]],s)
    s = re.sub(r'\\upcite\{([^}]+)\}',lambda m:r'\textsuperscript{['+','.join(k[1:] for k in m[1].split(','))+']}',s)
    def equation(m):
        expression=m[1]; label=re.search(r'\\label\{(eq:[^}]+)\}',expression)
        expression=re.sub(r'\\label\{[^}]+\}', '', expression)
        if label: expression += r'\qquad('+labels[label[1]]+')'
        return '$$'+expression.strip()+'$$'
    s = re.sub(r'\\begin\{equation\}(.*?)\\end\{equation\}',equation,s,flags=re.S)
    s = re.sub(r'\\label\{[^}]+\}','',s)
    s = re.sub(r'\\algtitle\{([^}]+)\}',lambda m:r'\paragraph{'+m[1]+'}',s)
    s = re.sub(r'\\begin\{tabular\*\}\{[^}]+\}\{@\{\\extracolsep\{\\fill\}\}([^}]+)\}',lambda m:r'\begin{tabular}{'+m[1]+'}',s)
    s = s.replace('\\end{tabular*}','\\end{tabular}')
    s = s.replace('figure*','figure').replace('table*','table')
    if name=='main':
        section=[-1]; subsection=[0]
        def heading(m):
            if m[1]=='section': section[0]+=1;subsection[0]=0; prefix=str(section[0])
            else: subsection[0]+=1;prefix=f'{section[0]}.{subsection[0]}'
            return '\\'+m[1]+'{'+prefix+' '+m[2]+'}'
        s = re.sub(r'\\(section|subsection)\{([^}]+)\}',heading,s)
    if '\\begin{thebibliography}' in s:
        a=s.index('\\begin{thebibliography}');b=s.index('\\end{thebibliography}')+len('\\end{thebibliography}')
        entries=re.findall(r'\\bibitem\{r(\d+)\}\s*(.*?)(?=\\bibitem|\\end\{thebibliography\})',s[a:b],re.S)
        s=s[:a]+'\\section*{参考文献}\n\n'+'\n\n'.join('['+i+'] '+text.strip() for i,text in entries)+s[b:]
    s = re.sub(r'figures/([^}]+)\.pdf',r'word-figures/\1.png',s)
    (QA/(name+'-word-input.tex')).write_text(s,encoding='utf-8',newline='\n')
    return QA/(name+'-word-input.tex')

def word():
    previews=OUT/'word-figures';previews.mkdir(exist_ok=True)
    for source in (OUT/'figures').glob('*.pdf'):
        subprocess.run([str(POPPLER),'-r','300','-singlefile','-png',str(source),str(previews/source.stem)],check=True,capture_output=True)
    template=Document();sec=template.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.left_margin=sec.right_margin=Mm(20);sec.top_margin=sec.bottom_margin=Mm(18)
    for name,size in [('Normal',10.5),('Title',17),('Heading 1',12),('Heading 2',10.5),('Heading 3',10.5),('Heading 4',10.5),('Caption',9)]:
        st=template.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.font.italic=False;st.font.bold=name.startswith('Heading') or name=='Title'
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体' if name in ('Title','Heading 1','Heading 2') else '宋体')
        pf=st.paragraph_format;pf.space_after=Pt(0);pf.line_spacing=1.15;pf.widow_control=True
        if name=='Normal': pf.first_line_indent=Pt(21)
        elif name.startswith('Heading'):pf.space_before=Pt(7);pf.space_after=Pt(4);pf.keep_with_next=True
    template.save(QA/'reference.docx')
    checks={}
    for name in ('main','supplement'):
        source=prepare_tex(name);dest=OUT/(name+'.docx')
        run=subprocess.run([str(PANDOC),str(source),'-f','latex','-t','docx','--resource-path',str(OUT),
            '--reference-doc',str(QA/'reference.docx'),'-o',str(dest)],capture_output=True,text=True,encoding='utf-8')
        (QA/(name+'-pandoc.log')).write_text(run.stderr,encoding='utf-8')
        assert run.returncode==0,run.stderr
        assert not re.search(r'Could not convert TeX math|Could not fetch resource',run.stderr),run.stderr
        doc=Document(dest)
        for element in (doc.styles.element,doc.element):
            for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
        for table in doc.tables:
            props=table._tbl.tblPr
            borders=props.find(qn('w:tblBorders'))
            if borders is not None:props.remove(borders)
            borders=OxmlElement('w:tblBorders')
            for side in ('top','left','bottom','right','insideH','insideV'):
                node=OxmlElement('w:'+side);node.set(qn('w:val'),'single' if side in ('top','bottom') else 'nil')
                node.set(qn('w:sz'),'6');borders.append(node)
            props.append(borders)
            for row_index,row in enumerate(table.rows):
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.first_line_indent=Pt(0)
                        for r in p.runs:r.font.size=Pt(9)
                    if row_index==0:
                        tcpr=cell._tc.get_or_add_tcPr(); cb=OxmlElement('w:tcBorders');edge=OxmlElement('w:bottom')
                        edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'6');cb.append(edge);tcpr.append(cb)
                if row_index==0:
                    repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        for mr in doc.element.iter(qn('m:r')):
            rp=mr.find(qn('w:rPr'))
            if rp is None:rp=OxmlElement('w:rPr');mr.insert(0,rp)
            sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');rp.append(sz)
        for p in doc.paragraphs:
            if p._p.find('.//'+qn('w:drawing')) is not None:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
                p.paragraph_format.line_spacing=1.0
            if p.style.name in ('Caption','Image Caption','Table Caption'):
                p.paragraph_format.first_line_indent=Pt(0)
            if re.match(r'^[图表]S?\d+ ',p.text):
                p.style=doc.styles['Caption'];p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.keep_with_next=p.text.startswith('表')
                for caption_run in p.runs:caption_run.font.size=Pt(9)
        if name=='main':
            title=next(p for p in doc.paragraphs if p.text.strip())
            title.style=doc.styles['Title']
            title.paragraph_format.first_line_indent=Pt(0)
            for title_run in title.runs:title_run.font.size=Pt(17);title_run.font.bold=True
        doc.core_properties.author='';doc.core_properties.title='资源匹配下超图支付网络的服务可靠性' if name=='main' else '补充分析S1–S5'
        doc.save(dest)
        with zipfile.ZipFile(dest) as z:
            root=etree.fromstring(z.read('word/document.xml'));ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            text=''.join(root.xpath('//w:t/text()|//m:t/text()',namespaces=ns))
            display=len(root.xpath('//m:oMathPara',namespaces=ns));images=len(doc.inline_shapes)
            assert images==(3 if name=='main' else 8),(name,images)
            # 18 numbered equations plus two unnumbered display expressions.
            assert name!='main' or display==20,(name,display)
            assert '\\begin{' not in text and '\\upcite' not in text
            checks[name]=dict(editable_math_objects=len(root.xpath('//m:oMath',namespaces=ns)),display_equations=display,
                images=images,tables=len(doc.tables),bytes=dest.stat().st_size,sha256=sha(dest),conversion_warnings=run.stderr)
    write_json(OUT/'word-checks.json',checks)

if __name__=='__main__':
    import sys
    if sys.argv[1]=='revise':revise()
    elif sys.argv[1]=='word':word()
