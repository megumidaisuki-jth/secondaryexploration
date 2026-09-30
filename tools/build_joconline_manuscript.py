"""Create the Chinese journal manuscript from locked numerical display inputs.

Uses the bundled artifact Python; optional plotting packages are supplied through
PYTHONPATH. This builder performs no simulation, bootstrap, or raw-block reads.
"""
from __future__ import annotations

import hashlib
import json
import re
import argparse
import shutil
from fractions import Fraction
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'manuscript/joconline'
FIG = OUT / 'figures'
INPUT = ROOT / 'manuscript/generated/synthetic-v1'
DOCX = OUT / '超图支付网络服务可靠性_通信学报中文稿.docx'


def digest(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def inputs():
    manifest = json.loads((INPUT / 'manifest.json').read_text(encoding='utf-8'))
    for name in ('numerical-registry.json', 'table-1.json'):
        assert digest(INPUT / name) == manifest['files'][name]['sha256']
    registry = json.loads((INPUT / 'numerical-registry.json').read_text(encoding='utf-8'))
    table = json.loads((INPUT / 'table-1.json').read_text(encoding='utf-8'))
    assert len(registry['phase_intervals']) == 80
    assert len(registry['cross_phase_replication']) == 40
    assert len(table) == 8
    return registry, table


def figures(registry):
    import matplotlib as mpl
    mpl.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle
    from matplotlib.lines import Line2D
    from matplotlib import font_manager
    for font in ('C:/Windows/Fonts/STZHONGS.TTF', 'C:/Windows/Fonts/times.ttf'):
        font_manager.fontManager.addfont(font)
    mpl.rcParams.update({
        'font.family': ['Times New Roman', 'STZhongsong', 'sans-serif'],
        'font.size': 7.5, 'axes.labelsize': 7.5, 'axes.titlesize': 7.5,
        'xtick.labelsize': 7.5, 'ytick.labelsize': 7.5, 'legend.fontsize': 7.5,
        'axes.unicode_minus': False, 'svg.fonttype': 'none', 'pdf.fonttype': 42,
        'axes.linewidth': .57, 'xtick.direction': 'in', 'ytick.direction': 'in',
        'svg.hashsalt': 'joconline-zh-v1',
    })
    FIG.mkdir(parents=True, exist_ok=True)
    def save(fig, stem):
        fig.savefig(FIG / f'{stem}.svg', metadata={'Date': None})
        fig.savefig(FIG / f'{stem}.pdf', metadata={'CreationDate': None, 'ModDate': None})
        fig.savefig(FIG / f'{stem}.png', dpi=600)
        plt.close(fig)

    fig, ax = plt.subplots(figsize=(170/25.4, 62/25.4))
    fig.subplots_adjust(left=.01, right=.99, bottom=.025, top=.975)
    ax.set(xlim=(0,1), ylim=(0,1)); ax.axis('off')
    def box(x,y,w,h,title,detail):
        ax.add_patch(Rectangle((x,y),w,h,facecolor='white',edgecolor='black',linewidth=.57))
        ax.text(x+w/2,y+h-.055,title,ha='center',va='top')
        ax.text(x+w/2,y+.08,detail,ha='center',va='bottom',linespacing=1.3)
    box(.015,.63,.28,.34,'独立父图与构造','3类模型 × 4种规模\n每规模每相位60个父图')
    box(.36,.63,.28,.34,'资源匹配与配对请求','4类超图及各自二元参照\n每父图7条留出轨迹')
    box(.705,.63,.28,.34,'服务终点','固定时域失败风险\n受限无路径时间 / H')
    box(.13,.14,.32,.30,'正式相位','240个区块 · 40项比较')
    box(.55,.14,.32,.30,'独立确认相位','240个区块 · 40项比较')
    for start,end in [((.295,.8),(.36,.8)),((.64,.8),(.705,.8)),((.29,.54),(.29,.44)),((.71,.54),(.71,.44))]:
        ax.annotate('',xy=end,xytext=start,arrowprops={'arrowstyle':'->','lw':.85,'color':'black'})
    ax.plot([.29,.845,.845],[.54,.54,.63],color='black',lw=.57)
    ax.text(.5,.56,'相同冻结协议 · 相位分别推断',ha='center',va='bottom')
    ax.text(.5,.03,'父图分层重抽样 → 区间与总体门判定 → 全部40项跨相位复现记录',ha='center',va='center')
    save(fig,'figure-1-design')

    rows={(r['phase'],r['contrast_id']):r for r in registry['phase_intervals']}
    metrics=('failure_risk','normalized_restricted_tau_nopath')
    sizes=(30,60,120,240)
    fig,axs=plt.subplots(1,2,figsize=(170/25.4,78/25.4))
    fig.subplots_adjust(left=.08,right=.985,bottom=.31,top=.96,wspace=.29)
    for p,(ax,metric) in enumerate(zip(axs,metrics)):
        for j,phase in enumerate(('formal','confirmation')):
            for i,n in enumerate(sizes):
                r=rows[(phase,f'{metric}.n{n:04d}.global')]
                lo,hi,est=(float(Fraction(*r[k])) for k in ('lower','upper','estimate'))
                y=i+(-.13 if j==0 else .13)
                ax.plot([lo,hi],[y,y],color='black',lw=.85,ls='-' if j==0 else '--')
                ax.plot(est,y,marker='o' if j==0 else 's',ms=4,mec='black',mfc='black' if j==0 else 'white',mew=.57)
        ax.axvline(0,color='0.45',ls=':',lw=.57)
        ax.set_ylim(3.5,-.5)
        ax.set_yticks(range(4),[str(n) for n in sizes])
        ax.set_ylabel('节点数')
        ax.set_xlim((-.9,.05) if p==0 else (-.03,.65))
        ax.set_xlabel('超图 − 资源匹配二元参照')
        ax.xaxis.set_major_locator(mpl.ticker.MaxNLocator(5))
        ax.text(.5,-.29,'(a) 固定时域失败风险差' if p==0 else '(b) 归一化受限无路径时间差',transform=ax.transAxes,ha='center')
    leg=fig.legend([Line2D([],[],c='black',marker='o',lw=.85),Line2D([],[],c='black',marker='s',mfc='white',ls='--',lw=.85)],['正式相位','确认相位'],ncol=2,loc='lower center',bbox_to_anchor=(.52,.01),frameon=True,fancybox=False,edgecolor='black')
    leg.get_frame().set_linewidth(.57)
    save(fig,'figure-2-global-contrasts')

    families=('demand-aware','fhs3','fhs5','global','nch')
    reps={r['contrast_id']:r for r in registry['cross_phase_replication']}
    labels={'independently_confirmed':'双相位','formal_only':'仅正式','confirmation_only':'仅确认','neither':'均未通过'}
    fig,ax=plt.subplots(figsize=(170/25.4,77/25.4))
    fig.subplots_adjust(left=.25,right=.985,bottom=.15,top=.86)
    for i,(metric,n) in enumerate((m,n) for m in metrics for n in sizes):
        for j,family in enumerate(families):
            r=reps[f'{metric}.n{n:04d}.{family}']
            ax.add_patch(Rectangle((j-.5,i-.5),1,1,facecolor='0.95',edgecolor='black',lw=.57))
            ax.text(j,i,labels[r['replication_state']],ha='center',va='center')
    ax.set(xlim=(-.5,4.5),ylim=(7.5,-.5))
    ax.set_xticks(range(5),['需求感知','FHS3','FHS5','总体','NCH'])
    ax.xaxis.tick_top(); ax.tick_params(length=0,pad=7)
    ax.set_yticks(range(8),[f'{"失败风险" if m==metrics[0] else "受限时间/H"} · {n}' for m in metrics for n in sizes])
    fig.text(.5,.055,'每格保留一个预定比较；构造族比较须在两个相位均满足总体门条件。',ha='center')
    save(fig,'figure-3-replication-matrix')


def create_docx(table, descriptive=None):
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from lxml import etree
    from latex2mathml.converter import convert

    source=(OUT/'manuscript-zh.md').read_text(encoding='utf-8')
    blocks=source.strip().split('\n\n')
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Mm(210); sec.page_height=Mm(297)
    sec.left_margin=sec.right_margin=Mm(20)
    sec.top_margin=Mm(18); sec.bottom_margin=Mm(18)
    sec.footer_distance=Mm(8)
    def set_font(style,size,cn='宋体',bold=False):
        style.font.name='Times New Roman'; style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(0,0,0); style.font.bold=bold
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),cn)
    set_font(doc.styles['Normal'],10.5)
    normal=doc.styles['Normal'].paragraph_format
    normal.space_after=Pt(0); normal.line_spacing=Pt(15)
    normal.first_line_indent=Pt(21)
    normal.widow_control=True
    set_font(doc.styles['Title'],18,'黑体')
    set_font(doc.styles['Heading 1'],12,'黑体')
    set_font(doc.styles['Heading 2'],10.5,'黑体')
    set_font(doc.styles['Caption'],9,'宋体')
    for name in ('Heading 1','Heading 2'):
        p=doc.styles[name].paragraph_format
        p.first_line_indent=Pt(0); p.space_before=Pt(7);p.space_after=Pt(4)
        p.keep_with_next=True;p.line_spacing=1.1
    cp=doc.styles['Caption'].paragraph_format
    cp.first_line_indent=Pt(0);cp.space_before=Pt(3);cp.space_after=Pt(5);cp.line_spacing=1.05
    doc.core_properties.title='资源匹配下超图支付网络的服务可靠性'
    doc.core_properties.author=''
    doc.core_properties.subject='通信学报中文论文稿'
    doc.core_properties.comments=''
    for root in (doc.styles.element, doc.element):
        for border in list(root.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    # Word's installed, tested MathML-to-OMML stylesheet creates editable math.
    transform=etree.XSLT(etree.parse('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'))
    eq_count=0; fig_count=0; body=False; refs=False
    inline_math={
        '𝓗=(V,𝓔)':r'\mathcal H=(V,\mathcal E)', 'n=|V|':r'n=|V|',
        'e∈𝓔':r'e\in\mathcal E','v∈e':r'v\in e',
        'xₑ,ᵥ(t)':r'x_{e,v}(t)','Cₑ':r'C_e','Bᵥ':r'B_v',
        'qₜ=(sₜ,dₜ,aₜ)':r'q_t=(s_t,d_t,a_t)',
        'P=(v₀,e₁,v₁,…,eℓ,vℓ)':r'P=(v_0,e_1,v_1,\ldots,e_\ell,v_\ell)',
        'v₀=sₜ':r'v_0=s_t','vℓ=dₜ':r'v_\ell=d_t',
        'vⱼ₋₁,vⱼ∈eⱼ':r'v_{j-1},v_j\in e_j',
        'xₑⱼ,ᵥⱼ₋₁(t−1)≥aₜ':r'x_{e_j,v_{j-1}}(t-1)\geq a_t',
        '𝓟ₜ':r'\mathcal P_t','τdep':r'\tau_{\mathrm{dep}}',
        'τnp':r'\tau_{\mathrm{np}}','τrej':r'\tau_{\mathrm{rej}}',
        'G=(V,E)':r'G=(V,E)','Dᵤᵥ':r'D_{uv}','mᵤᵥ':r'm_{uv}',
        'Kf':r'K_f','Z∈{Y,F}':r'Z\in\{Y,F\}',
    }
    def math_node(expr):
        node=transform(etree.fromstring(convert(expr).encode())).getroot()
        for mr in node.iter(qn('m:r')):
            rp=OxmlElement('w:rPr')
            sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');rp.append(sz)
            mr.insert(0,rp)
        return node
    def text_runs(p,text,sup=True):
        if body and sup:
            for literal,expr in sorted(inline_math.items(),key=lambda kv:-len(kv[0])):
                text=text.replace(literal,'$'+expr+'$')
        for chunk in re.split(r'(\$[^$]+\$|\[\d+(?:[-,]\d+)*\])',text):
            if chunk.startswith('$') and chunk.endswith('$'):
                p._p.append(math_node(chunk[1:-1]));continue
            run=p.add_run(chunk)
            if sup and re.fullmatch(r'\[\d+(?:[-,]\d+)*\]',chunk):
                run.font.superscript=True
    def render_table(caption, heads, rows, widths):
        cap=doc.add_paragraph(caption,'Caption')
        cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.keep_with_next=True
        t=doc.add_table(rows=1,cols=len(heads))
        t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for c,w in zip(t.columns,widths): c.width=Mm(w)
        for c,txt in zip(t.rows[0].cells,heads): c.text=txt
        for vals in rows:
            cells=t.add_row().cells
            for c,txt,w in zip(cells,vals,widths): c.text=txt;c.width=Mm(w)
        for ri,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr()
            cant=OxmlElement('w:cantSplit');trpr.append(cant)
            if ri==0: trpr.append(OxmlElement('w:tblHeader'))
            for ci,c in enumerate(row.cells):
                c.width=Mm(widths[ci])
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=c._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right','insideH','insideV'):
                    el=OxmlElement('w:'+edge)
                    visible=(ri==0 and edge in ('top','bottom')) or (ri==len(t.rows)-1 and edge=='bottom')
                    el.set(qn('w:val'),'single' if visible else 'nil')
                    if visible: el.set(qn('w:sz'),'6');el.set(qn('w:color'),'000000')
                    borders.append(el)
                pr.append(borders)
                for p in c.paragraphs:
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.keep_with_next=ri < len(t.rows)-1
                    p.paragraph_format.first_line_indent=Pt(0)
                    p.paragraph_format.line_spacing=Pt(13)
                    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                    for r in p.runs: r.font.size=Pt(9);r.bold=(ri==0)
    def global_table():
        rows=[]
        for row in table:
            vals=['失败风险' if row['formal']['metric']=='failure_risk' else '受限无路径时间/H',str(row['formal']['node_count'])]
            for phase in ('formal','confirmation'):
                d=row[phase]['display']; vals.append(f"{d['estimate']} [{d['lower']}, {d['upper']}]")
            rows.append(vals)
        render_table('表1　总体配对效应及多重比较校正区间',
                     ('终点','节点数','正式相位估计［区间］','确认相位估计［区间］'),rows,(31,17,61,61))
    def model_table():
        rows=[]
        names={'barabasi_albert':'BA','er_gnm':'ER-GNM','sbm_fixed_count':'固定边数SBM'}
        for n in (30,60,120,240):
            for model,label in names.items():
                row=[str(n),label]
                for phase in ('formal','confirmation'):
                    r=next(r for r in descriptive['phases'][phase]['parent_stratum_summaries']
                           if r['metric']=='failure_risk' and r['source_family']=='global'
                           and r['node_count']==n and r['parent_model']==model)
                    assert r['count']==20
                    scaled=round(Fraction(*r['mean'])*1000)
                    row.append(('-' if scaled<0 else '')+f'{abs(scaled)//1000}.{abs(scaled)%1000:03d}')
                rows.append(row)
        render_table('表2　总体失败风险差的父图模型分层均值',
                     ('节点数','父图模型','正式相位','确认相位'),rows,(22,48,50,50))
        doc.add_paragraph('注：每格为20个独立父图的描述性均值，不是新增的模型间显著性检验。','Caption')
    def activity_table():
        rows=[]
        for n in (30,60,120,240):
            row=[str(n)]
            for phase in ('formal','confirmation'):
                counts=[]
                for activity in ('changed','unchanged'):
                    r=next(r for r in descriptive['phases'][phase]['activity_sensitivity']
                           if r['metric']=='failure_risk' and r['node_count']==n
                           and r['topology_activity']==activity)
                    counts.append(sum(s['parent_count'] for s in r['strata']))
                assert sum(counts)==60
                row.extend(map(str,counts))
            rows.append(row)
        render_table('表3　需求感知搜索的拓扑改变与未改变父图数量',
                     ('节点数','正式改变','正式未改变','确认改变','确认未改变'),rows,(22,37,37,37,37))
        doc.add_paragraph('注：每个规模、每个相位共60个父图。改变以最终拓扑相对于FHS5种子是否变化定义，不表示服务改善。','Caption')
    for block in blocks:
        if block.startswith('# '):
            p=doc.add_paragraph(block[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(8)
        elif block.startswith('## Service '):
            p=doc.add_paragraph(block[3:]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(4)
            for r in p.runs:r.font.size=Pt(13);r.bold=True
        elif block.startswith('## '):
            body=True
            refs=block=='## 参考文献'
            doc.add_paragraph(block[3:],'Heading 1')
        elif block.startswith('### '):
            doc.add_paragraph(block[4:],'Heading 2')
        elif block.startswith('$$'):
            eq_count+=1
            if doc.paragraphs:
                doc.paragraphs[-1].paragraph_format.keep_with_next=True
            expr=block.strip('$ ').strip()
            match=re.search(r'\\tag\{(\d+)\}',expr)
            num=match[1];expr=re.sub(r'\\tag\{\d+\}','',expr).strip()
            omml=math_node(expr)
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(82),WD_TAB_ALIGNMENT.CENTER)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(169),WD_TAB_ALIGNMENT.RIGHT)
            p.add_run('\t');p._p.append(omml);p.add_run('\t('+num+')')
        elif block.startswith('!['):
            file=re.search(r'\]\(([^)]+)\)',block)[1]
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.line_spacing=1.0
            p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(OUT/file),width=Mm(170));fig_count+=1
        elif block=='{{TABLE_GLOBAL}}':
            global_table()
        elif block=='{{TABLE_MODEL}}':
            assert descriptive is not None
            model_table()
        elif block=='{{TABLE_ACTIVITY}}':
            assert descriptive is not None
            activity_table()
        elif block.startswith('算法'):
            p=doc.add_paragraph(block)
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_before=Pt(5)
            for r in p.runs:r.bold=True
        elif re.match(r'^[图表]\d+[ \u3000]',block):
            if block.startswith('表1'):
                block='注：'+block.split('。',1)[1]
            p=doc.add_paragraph(block,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            p=doc.add_paragraph()
            if refs and re.match(r'^\[\d+\]',block):
                text_runs(p,block,False)
                p.paragraph_format.first_line_indent=Pt(-18);p.paragraph_format.left_indent=Pt(18)
                p.paragraph_format.line_spacing=Pt(12);p.paragraph_format.space_after=Pt(3)
                for r in p.runs:r.font.size=Pt(9)
            else:
                text_runs(p,block)
                p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            if not body:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(12)
                p.paragraph_format.space_after=Pt(4)
                if block.startswith(('作者姓名','（作者单位','AUTHOR ','Affiliations')):
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(9)
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Pt(0)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    doc.save(DOCX)
    abstract=next(b for b in blocks if b.startswith('摘要：'))[3:]
    title=blocks[0][2:]
    assert len(abstract)<=200,(len(abstract),'abstract too long')
    assert len(title)<=20
    assert eq_count==(16 if descriptive is not None else 9) and fig_count==3
    cited=[]
    content=source.split('## 参考文献')[0]
    for hit in re.findall(r'\[(\d+(?:[-,]\d+)*)\]',content):
        for piece in hit.split(','):
            if '-' in piece:
                a,b=map(int,piece.split('-'));values=range(a,b+1)
            else: values=[int(piece)]
            for value in values:
                if value not in cited:cited.append(value)
    assert cited==list(range(1,16)),cited
    return {'chinese_title_characters':len(title),'chinese_abstract_characters':len(abstract),
            'native_display_equations':eq_count,'figures':fig_count,'tables':3 if descriptive is not None else 1,'references':len(cited)}


def main():
    global OUT,FIG,DOCX
    parser=argparse.ArgumentParser()
    parser.add_argument('--expanded',action='store_true')
    args=parser.parse_args()
    descriptive=None
    if args.expanded:
        OUT=ROOT/'manuscript/joconline-expanded'
        FIG=OUT/'figures'
        DOCX=OUT/'超图支付网络服务可靠性_通信学报中文扩展稿.docx'
        folder=ROOT/'manuscript/generated/results-writing-v1'
        source_manifest=json.loads((folder/'manifest.json').read_text(encoding='utf-8'))
        assert digest(folder/'evidence-summary.json')==source_manifest['files']['evidence-summary.json']
        descriptive=json.loads((folder/'evidence-summary.json').read_text(encoding='utf-8'))
    registry,table=inputs()
    if args.expanded:
        old=ROOT/'manuscript/joconline'
        original=json.loads((old/'manifest.json').read_text(encoding='utf-8'))
        FIG.mkdir(parents=True,exist_ok=True)
        for p in (old/'figures').iterdir():
            assert digest(p)==original['files'][str(p.relative_to(old))]['sha256']
            shutil.copy2(p,FIG/p.name)
    else:
        figures(registry)
    counts=create_docx(table,descriptive)
    files=[DOCX,OUT/'manuscript-zh.md',*sorted(FIG.glob('*'))]
    manifest={'schema':'joconline-manuscript.v1','source_registry_sha256':digest(INPUT/'numerical-registry.json'),
              'source_table_sha256':digest(INPUT/'table-1.json'),'builder_sha256':digest(__file__),
              'counts':counts,'scientific_input_rows':{'intervals':80,'replications':40,'global_comparisons':8},
              'files':{str(p.relative_to(OUT)):{'bytes':p.stat().st_size,'sha256':digest(p)} for p in files}}
    if args.expanded:
        manifest['source_descriptive_sha256']=digest(folder/'evidence-summary.json')
        manifest['previous_version_commit']='045d5f80530a57875cf33595b0cb69b8569d8a3d'
    (OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'output':str(DOCX),**counts},ensure_ascii=False))


if __name__=='__main__':
    main()
